"""Load the DVAGO spreadsheet into SQLite.

    uv run python scripts/build_dvago_db.py [--xlsx PATH] [--db PATH]

Thin on purpose. All this does is turn a workbook into ``{table: [row, ...]}``;
the validation and the build live in ``agent.data.loader``, where they are
testable without a spreadsheet and where the running agent can reach them. That
also keeps ``openpyxl`` where it belongs — a development dependency, never
imported by the agent.
"""

from __future__ import annotations

import argparse
import sys
from pathlib import Path
from typing import Any

sys.path.insert(0, str(Path(__file__).resolve().parents[1] / "src"))

from agent.data.loader import ValidationError, write_database  # noqa: E402
from agent.data.schema_meta import load_schema  # noqa: E402

DEFAULT_XLSX = Path("data/dvago/DVAGO_demo_dataset.xlsx")
DEFAULT_DB = Path("data/dvago/dvago.db")

# Sheets that document the workbook rather than hold a table.
META_SHEETS = {"_README", "_data_dictionary"}


def read_workbook(path: Path) -> dict[str, list[dict[str, Any]]]:
    from openpyxl import load_workbook

    workbook = load_workbook(path, read_only=True, data_only=True)
    tables: dict[str, list[dict[str, Any]]] = {}
    try:
        for name in workbook.sheetnames:
            if name in META_SHEETS:
                continue
            rows = workbook[name].iter_rows(values_only=True)
            try:
                header = [str(cell) for cell in next(rows) if cell is not None]
            except StopIteration:
                tables[name] = []
                continue
            tables[name] = [
                dict(zip(header, row, strict=False))
                for row in rows
                if any(cell is not None for cell in row)
            ]
    finally:
        workbook.close()
    return tables


def main() -> int:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--xlsx", type=Path, default=DEFAULT_XLSX)
    parser.add_argument("--db", type=Path, default=DEFAULT_DB)
    args = parser.parse_args()

    if not args.xlsx.exists():
        print(f"no workbook at {args.xlsx}")
        print("run: uv run python scripts/generate_dvago_dataset.py")
        return 1

    try:
        counts = write_database(read_workbook(args.xlsx), load_schema(), args.db)
    except ValidationError as error:
        print(f"{args.xlsx} failed validation - database not written\n")
        for problem in error.problems[:40]:
            print(f"  - {problem}")
        if len(error.problems) > 40:
            print(f"  ... and {len(error.problems) - 40} more")
        return 1

    size_kb = args.db.stat().st_size / 1024
    print(f"wrote {args.db}  ({sum(counts.values()):,} rows, {size_kb:,.0f} KB)")
    for name, count in counts.items():
        print(f"  {name:24} {count:>6,}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
