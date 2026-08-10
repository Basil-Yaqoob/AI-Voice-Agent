"""The generated dataset, when it exists.

Skipped on a fresh clone — the workbook and database are build artifacts, not
source. When they are present these check the properties a demo depends on:
that the anchors are where the runbook says they are, that the spreadsheet and
the DDL still agree, and that the referential integrity survived the round trip
through Excel.

The anchor checks are the ones that matter. A demo script is written against
specific rows, and a regenerate that quietly moves them turns a rehearsed call
into an improvised one.
"""

from __future__ import annotations

import sqlite3
from collections.abc import Iterator
from pathlib import Path

import pytest

from agent.data.schema_meta import load_schema

PROJECT_ROOT = Path(__file__).resolve().parents[2]
DB_PATH = PROJECT_ROOT / "data" / "dvago" / "dvago.db"
XLSX_PATH = PROJECT_ROOT / "data" / "dvago" / "DVAGO_demo_dataset.xlsx"

pytestmark = pytest.mark.skipif(
    not DB_PATH.exists(),
    reason="dataset not built; run scripts/generate_dvago_dataset.py then build_dvago_db.py",
)


@pytest.fixture(scope="module")
def conn() -> Iterator[sqlite3.Connection]:
    connection = sqlite3.connect(f"file:{DB_PATH}?mode=ro", uri=True)
    connection.row_factory = sqlite3.Row
    yield connection
    connection.close()


def _one(conn: sqlite3.Connection, sql: str, *params: object) -> sqlite3.Row | None:
    return conn.execute(sql, params).fetchone()


# ── integrity ────────────────────────────────────────────────────────────


def test_no_foreign_key_violations(conn: sqlite3.Connection) -> None:
    assert conn.execute("PRAGMA foreign_key_check").fetchall() == []


def test_every_table_has_rows(conn: sqlite3.Connection) -> None:
    empty = [
        name
        for name in load_schema()
        if conn.execute(f"SELECT count(*) FROM {name}").fetchone()[0] == 0
    ]
    assert empty == []


def test_the_dataset_declares_itself_synthetic(conn: sqlite3.Connection) -> None:
    """Fabricated records about a real company must never look like real ones."""
    row = _one(conn, "SELECT value FROM dataset_meta WHERE key = 'synthetic'")
    assert row is not None
    assert row["value"] == "yes"


def test_order_totals_add_up(conn: sqlite3.Connection) -> None:
    wrong = conn.execute(
        """SELECT order_id FROM orders
           WHERE abs(total_pkr - (subtotal_pkr - discount_pkr + delivery_fee_pkr)) > 0.01"""
    ).fetchall()
    assert wrong == []


def test_line_totals_add_up(conn: sqlite3.Connection) -> None:
    wrong = conn.execute(
        "SELECT order_item_id FROM order_items "
        "WHERE abs(line_total_pkr - quantity * unit_price_pkr) > 0.01"
    ).fetchall()
    assert wrong == []


def test_reserved_never_exceeds_stock(conn: sqlite3.Connection) -> None:
    assert conn.execute(
        "SELECT count(*) FROM inventory WHERE qty_reserved > qty_on_hand"
    ).fetchone()[0] == 0


def test_cold_chain_items_only_sit_where_there_is_a_fridge(
    conn: sqlite3.Connection,
) -> None:
    misplaced = conn.execute(
        """SELECT i.branch_id, i.sku FROM inventory i
           JOIN products p ON p.sku = i.sku
           JOIN branches b ON b.branch_id = i.branch_id
           WHERE p.is_cold_chain = 1 AND b.has_cold_chain = 0 AND i.qty_on_hand > 0"""
    ).fetchall()
    assert misplaced == []


def test_delivered_orders_have_a_delivery_time(conn: sqlite3.Connection) -> None:
    assert conn.execute(
        "SELECT count(*) FROM orders WHERE status = 'delivered' AND delivered_at IS NULL"
    ).fetchone()[0] == 0


def test_no_stock_is_already_expired(conn: sqlite3.Connection) -> None:
    row = _one(conn, "SELECT value FROM dataset_meta WHERE key = 'as_of'")
    assert row is not None
    stale = conn.execute(
        "SELECT count(*) FROM inventory WHERE expiry_date < ? AND qty_on_hand > 0",
        (row["value"],),
    ).fetchone()[0]
    assert stale == 0


def test_phone_numbers_look_pakistani(conn: sqlite3.Connection) -> None:
    bad = conn.execute(
        "SELECT phone FROM customers WHERE phone NOT LIKE '03%' OR length(phone) != 11"
    ).fetchall()
    assert bad == []


# ── anchors the demo script depends on ───────────────────────────────────


def test_the_main_caller_exists_with_the_allergy(conn: sqlite3.Connection) -> None:
    row = _one(conn, "SELECT * FROM customers WHERE phone = '03001234567'")
    assert row is not None
    assert row["full_name"] == "Ayesha Khan"
    assert row["loyalty_tier"] == "gold"
    assert row["allergies"] == "penicillin"
    assert row["full_name_ur"]  # the greeting needs the Urdu spelling


def test_the_unknown_caller_number_stays_unknown(conn: sqlite3.Connection) -> None:
    """03009999999 demonstrates the new-customer path and must never be assigned."""
    assert _one(conn, "SELECT 1 FROM customers WHERE phone = '03009999999'") is None


def test_the_live_order_is_out_for_delivery(conn: sqlite3.Connection) -> None:
    row = _one(
        conn,
        """SELECT o.status, d.eta, d.rider_id FROM orders o
           JOIN deliveries d ON d.order_id = o.order_id
           WHERE o.order_id = 'DV-2026-104471'""",
    )
    assert row is not None
    assert row["status"] == "out_for_delivery"
    assert row["eta"]
    assert row["rider_id"]


def test_augmentin_is_out_at_bahadurabad_but_in_stock_nearby(
    conn: sqlite3.Connection,
) -> None:
    """The demo that matters: not 'no', but 'not here'."""
    home = _one(
        conn,
        """SELECT qty_available FROM v_available_stock
           WHERE area = 'Bahadurabad'
             AND sku = (SELECT sku FROM products WHERE brand_name = 'Augmentin')""",
    )
    assert home is not None
    assert home["qty_available"] == 0

    nearby = conn.execute(
        """SELECT count(*) FROM v_available_stock
           WHERE city = 'Karachi' AND qty_available > 0
             AND sku = (SELECT sku FROM products WHERE brand_name = 'Augmentin')"""
    ).fetchone()[0]
    assert nearby > 0


def test_the_substitute_path_is_reachable(conn: sqlite3.Connection) -> None:
    """Nan Pro must be out across all of Karachi, or the agent offers a branch instead."""
    available = conn.execute(
        """SELECT count(*) FROM v_available_stock
           WHERE city = 'Karachi' AND qty_available > 0
             AND sku = (SELECT sku FROM products WHERE brand_name = 'Nan Pro 1')"""
    ).fetchone()[0]
    assert available == 0

    alternatives = conn.execute(
        """SELECT count(*) FROM substitutes
           WHERE sku = (SELECT sku FROM products WHERE brand_name = 'Nan Pro 1')"""
    ).fetchone()[0]
    assert alternatives >= 1


def test_the_prescription_hold_is_pending(conn: sqlite3.Connection) -> None:
    row = _one(
        conn,
        """SELECT rx.status FROM orders o
           JOIN prescriptions rx ON rx.prescription_id = o.prescription_id
           WHERE o.order_id = 'DV-2026-104502'""",
    )
    assert row is not None
    assert row["status"] == "pending_review"


def test_the_refund_is_still_in_flight(conn: sqlite3.Connection) -> None:
    row = _one(conn, "SELECT status FROM refunds WHERE refund_id = 'REF-2026-0071'")
    assert row is not None
    assert row["status"] == "processing"


def test_the_open_complaint_is_open(conn: sqlite3.Connection) -> None:
    row = _one(conn, "SELECT status FROM complaints WHERE complaint_id = 'CMP-2026-0042'")
    assert row is not None
    assert row["status"] in {"open", "in_progress"}


def test_there_is_an_expired_promo_and_a_live_one(conn: sqlite3.Connection) -> None:
    assert _one(conn, "SELECT 1 FROM promotions WHERE promo_code = 'HBL20' AND is_active = 0")
    assert _one(conn, "SELECT 1 FROM promotions WHERE promo_code = 'DVAGO10' AND is_active = 1")


def test_a_courier_only_city_customer_exists(conn: sqlite3.Connection) -> None:
    """Someone the agent must not promise one-hour delivery to."""
    row = _one(conn, "SELECT city FROM customers WHERE phone = '03451239876'")
    assert row is not None
    assert row["city"] not in {
        r["city"] for r in conn.execute("SELECT DISTINCT city FROM branches")
    }


# ── the workbook ─────────────────────────────────────────────────────────


@pytest.mark.skipif(not XLSX_PATH.exists(), reason="workbook not generated")
def test_workbook_sheets_match_the_schema() -> None:
    from openpyxl import load_workbook

    workbook = load_workbook(XLSX_PATH, read_only=True)
    try:
        sheets = set(workbook.sheetnames)
    finally:
        workbook.close()
    assert set(load_schema()) <= sheets
    assert {"_README", "_data_dictionary"} <= sheets


@pytest.mark.skipif(not XLSX_PATH.exists(), reason="workbook not generated")
def test_workbook_headers_match_the_schema_exactly() -> None:
    """The loader writes columns positionally. A reordered header silently corrupts."""
    from openpyxl import load_workbook

    workbook = load_workbook(XLSX_PATH, read_only=True)
    try:
        for name, table in load_schema().items():
            header = [c.value for c in next(workbook[name].iter_rows(max_row=1))]
            assert header == list(table.column_names), name
    finally:
        workbook.close()


@pytest.mark.skipif(not XLSX_PATH.exists(), reason="workbook not generated")
def test_data_dictionary_covers_every_column() -> None:
    from openpyxl import load_workbook

    workbook = load_workbook(XLSX_PATH, read_only=True)
    try:
        rows = list(workbook["_data_dictionary"].iter_rows(min_row=2, values_only=True))
    finally:
        workbook.close()

    documented = {(row[0], row[1]) for row in rows}
    expected = {
        (table.name, column.name)
        for table in load_schema().values()
        for column in table.columns
    }
    assert expected == documented
    assert all(row[8] for row in rows), "every column needs a description"
